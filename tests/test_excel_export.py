"""勤務表Excel書き出し(excel_export.py)の集計・並び替えロジックのテスト。

SUMMARY の「残業」列は実物の勤務表にある手書き欄なので、
自動では埋めない(常に空欄)仕様。ここではテストしない。
"""

from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent / "src"))

from shift_scheduler.excel_export import (  # noqa: E402
    SUMMARY,
    SUMMARY_MARKS,
    _countif_formula,
    _order_staff,
    _split_mark,
)
from shift_scheduler.models import StaffProfile  # noqa: E402
from shift_scheduler.scheduler import ScheduleResult  # noqa: E402


# --- 集計の式 ------------------------------------------------------------------


def test_集計は数字ではなく式で入れる():
    """受け取った側が手直ししたら、その場で合計が変わる必要がある。"""
    formula = _countif_formula("E6:K6", ("公",))
    assert formula == '=COUNTIF(E6:K6,"公")'


def test_複数の記号はすべて足し合わせる():
    formula = _countif_formula("E6:K6", ("研", "健"))
    assert formula == '=COUNTIF(E6:K6,"研")+COUNTIF(E6:K6,"健")'


def test_残業は空欄のまま():
    """実物の手書き欄なので、記号からは決められない。"""
    assert SUMMARY_MARKS["残業"] == ()
    assert _countif_formula("E6:K6", SUMMARY_MARKS["残業"]) is None


def test_日勤は番号付きの勤務だけを数える():
    """時間を直接書く方のセル(「9-16時」など)は日勤に数えない(確認済み)。"""
    assert SUMMARY_MARKS["日勤"] == ("日①", "日②", "日③", "日", "日⑤", "日⑦", "日⑨")
    formula = _countif_formula("E6:K6", SUMMARY_MARKS["日勤"])
    assert "9-16時" not in formula
    assert "09:00" not in formula


def test_集計の項目がすべて定義されている():
    assert set(SUMMARY) == set(SUMMARY_MARKS)


# --- _split_mark ----------------------------------------------------------------


def test_split_mark_splits_on_naka_nuke():
    """中抜け「09:00-12:00・19:00-20:00」は2つに分かれること。"""
    main, second = _split_mark("09:00-12:00・19:00-20:00")
    assert main == "09:00-12:00"
    assert second == "19:00-20:00"


def test_split_mark_leaves_ordinary_marks_untouched():
    for mark in ("日①", "○", "◉", "公", ""):
        main, second = _split_mark(mark)
        assert main == mark
        assert second == ""


# --- _order_staff ----------------------------------------------------------------


def _profile(**kwargs):
    defaults = dict(staff_id="", name="", sheet_label="")
    defaults.update(kwargs)
    return StaffProfile(**defaults)


def _result_with(profiles):
    return ScheduleResult(
        status="最適",
        year=2026,
        month=10,
        assignments={p.staff_id: {} for p in profiles},
    )


def test_order_staff_puts_head_nurse_first_then_units_then_support_staff():
    head = _profile(staff_id="head", name="師長", is_head_nurse=True)
    bara = _profile(staff_id="bara", name="A", unit="ばら", sheet_label="1")
    sakura = _profile(staff_id="sakura", name="B", unit="さくら", sheet_label="1")
    yuri = _profile(staff_id="yuri", name="C", unit="ゆり", sheet_label="1")
    sumire = _profile(staff_id="sumire", name="D", unit="すみれ", sheet_label="1")
    support = _profile(staff_id="support", name="E", is_support_staff=True, sheet_label="①")

    # わざと並びを崩して渡す
    profiles = [support, sumire, yuri, sakura, bara, head]
    result = _result_with(profiles)

    ordered = [p.staff_id for p in _order_staff(profiles, result)]
    assert ordered == ["head", "bara", "sakura", "yuri", "sumire", "support"]


def test_order_staff_mixes_plain_numbers_and_circled_numbers_correctly():
    """看護(1,2,3…)を先に、介護(①②③…)を後に並べる。丸数字と半角数字が混在してもよい。"""
    profiles = [
        _profile(staff_id="c2", name="F", unit="ばら", sheet_label="②"),
        _profile(staff_id="n2", name="G", unit="ばら", sheet_label="2"),
        _profile(staff_id="c1", name="H", unit="ばら", sheet_label="①"),
        _profile(staff_id="n1", name="I", unit="ばら", sheet_label="1"),
    ]
    result = _result_with(profiles)

    ordered = [p.staff_id for p in _order_staff(profiles, result)]
    assert ordered == ["n1", "n2", "c1", "c2"]


def test_order_staff_excludes_people_not_in_the_result():
    """勤務表に載らない人(休職中など)は除く。"""
    listed = _profile(staff_id="a", name="A", unit="ばら", sheet_label="1")
    not_listed = _profile(staff_id="b", name="B", unit="ばら", sheet_label="2")
    result = _result_with([listed])

    ordered = [p.staff_id for p in _order_staff([listed, not_listed], result)]
    assert ordered == ["a"]


# --- 書き出したExcelの式が正しい数を出すか ------------------------------------------


def _evaluate_countif(sheet, formula: str) -> int:
    """=COUNTIF(範囲,"記号")+... を、実際のセルの中身に当てて数える。

    openpyxl は式を計算しないので、ここで同じことをして突き合わせる。
    """
    total = 0
    for part in formula.lstrip("=").split("+"):
        inside = part[part.index("(") + 1 : part.rindex(")")]
        cell_range, _, mark = inside.partition(",")
        mark = mark.strip('"')
        for row in sheet[cell_range]:
            for cell in row:
                if cell.value == mark:
                    total += 1
    return total


def _build_sheet(tmp_path):
    from datetime import date

    from openpyxl import load_workbook

    from shift_scheduler.calendar_utils import Day
    from shift_scheduler.excel_export import export_schedule

    profiles = [
        StaffProfile(staff_id="a", name="職員A", role="看護師", sheet_label="1", unit="ばら"),
        StaffProfile(staff_id="b", name="職員B", role="介護士", sheet_label="①", unit="ばら"),
    ]
    days = [Day(date(2026, 9, day)) for day in range(1, 6)]
    result = ScheduleResult(
        status="最適",
        year=2026,
        month=9,
        days=days,
        assignments={
            "a": {1: "日①", 2: "日", 3: "公", 4: "有", 5: "○"},
            "b": {1: "◉", 2: "公", 3: "研", 4: "健", 5: "夏"},
        },
    )
    path = tmp_path / "out.xlsx"
    export_schedule(result, profiles, path)
    return load_workbook(path)[  # 式をそのまま読む
        "勤務計画表"
    ]


def test_右端の集計の式が正しい数を出す(tmp_path):
    sheet = _build_sheet(tmp_path)
    from shift_scheduler.excel_export import FIRST_DAY_COLUMN, FIRST_STAFF_ROW

    expected = {
        FIRST_STAFF_ROW: {"日勤": 2, "夜勤": 1, "深夜": 0, "公休": 1, "有休": 1, "夏正": 0, "出研": 0},
        FIRST_STAFF_ROW + 2: {"日勤": 0, "夜勤": 0, "深夜": 1, "公休": 1, "有休": 0, "夏正": 1, "出研": 2},
    }
    for row, wanted in expected.items():
        for offset, label in enumerate(SUMMARY):
            cell = sheet.cell(row=row, column=FIRST_DAY_COLUMN + 5 + offset)
            if label == "残業":
                assert cell.value is None, "残業は空欄のまま"
                continue
            assert str(cell.value).startswith("="), f"{label}は式であること"
            assert _evaluate_countif(sheet, cell.value) == wanted[label], (
                f"{row}行目の{label}: 式 {cell.value}"
            )


def test_日別チェックの式が正しい数を出す(tmp_path):
    sheet = _build_sheet(tmp_path)
    from shift_scheduler.excel_export import FIRST_DAY_COLUMN

    found = None
    for row in range(1, sheet.max_row + 1):
        if sheet.cell(row=row, column=4).value == "日①":
            found = row
            break
    assert found, "日別チェックの「日①」行が見つからない"

    cell = sheet.cell(row=found, column=FIRST_DAY_COLUMN)  # 1日
    assert str(cell.value).startswith("=COUNTIF(")
    assert _evaluate_countif(sheet, cell.value) == 1, "1日の日①は1人"

    cell = sheet.cell(row=found, column=FIRST_DAY_COLUMN + 1)  # 2日
    assert _evaluate_countif(sheet, cell.value) == 0, "2日の日①は0人"


def test_そろっているかの判定も式で入る(tmp_path):
    sheet = _build_sheet(tmp_path)
    for row in range(1, sheet.max_row + 1):
        if sheet.cell(row=row, column=2).value == "6種そろっているか":
            value = str(sheet.cell(row=row, column=5).value)
            assert value.startswith("=IF(AND("), value
            assert '"○","×"' in value
            return
    raise AssertionError("判定行が見つからない")


# --- 手直し用のプルダウン ---------------------------------------------------------


def test_勤務欄にプルダウンが付く(tmp_path):
    """◉ や 丸数字 は打ちにくく、○ は似た字と取り違えやすい。"""
    from openpyxl.utils import range_boundaries

    from shift_scheduler.excel_export import EDIT_CHOICES

    sheet = _build_sheet(tmp_path)
    validations = sheet.data_validations.dataValidation
    assert len(validations) == 1, "プルダウンは1つにまとめる"

    validation = validations[0]
    for mark in ("○", "△", "◉", "公", "日①", "有"):
        assert mark in validation.formula1, f"{mark} が選択肢に無い"

    ranges = str(validation.sqref).split()
    assert len(ranges) == 2, "スタッフ2名ぶんの行に付く"
    for rg in ranges:
        left, top, right, bottom = range_boundaries(rg)
        assert top == bottom, "1行ずつ付ける"
        assert top % 2 == 0, "1人2行のうち1行目(偶数行)だけ"


def test_プルダウン以外の入力も許す(tmp_path):
    """時間を直接書く方のセルには「15」「19-20」なども入る。

    入力を禁止すると、こちらが想定しない直し方ができなくなる。
    """
    sheet = _build_sheet(tmp_path)
    validation = sheet.data_validations.dataValidation[0]
    assert validation.showErrorMessage is False
    assert validation.allow_blank is True


def test_集計欄と2行目にはプルダウンを付けない(tmp_path):
    from openpyxl.utils import range_boundaries

    from shift_scheduler.excel_export import FIRST_DAY_COLUMN, FIRST_STAFF_ROW

    sheet = _build_sheet(tmp_path)
    covered = set()
    for validation in sheet.data_validations.dataValidation:
        for rg in str(validation.sqref).split():
            left, top, right, bottom = range_boundaries(rg)
            for row in range(top, bottom + 1):
                for column in range(left, right + 1):
                    covered.add((row, column))

    # 2行目(中抜け・せ の欄)
    assert not any(
        (FIRST_STAFF_ROW + 1, column) in covered
        for column in range(FIRST_DAY_COLUMN, FIRST_DAY_COLUMN + 5)
    )
    # 右端の集計欄(計算式)
    assert not any(
        (FIRST_STAFF_ROW, FIRST_DAY_COLUMN + 5 + offset) in covered
        for offset in range(len(SUMMARY))
    )


def test_選択肢が長すぎればプルダウンを付けない(monkeypatch, tmp_path):
    """Excelの上限(255文字)を超えるとファイルが壊れる。付けない方を選ぶ。"""
    from shift_scheduler import excel_export

    monkeypatch.setattr(excel_export, "EDIT_CHOICES", tuple(f"記号{i:03d}" for i in range(60)))
    sheet = _build_sheet(tmp_path)

    assert len(sheet.data_validations.dataValidation) == 0
    # 勤務表そのものは書き出せていること
    assert sheet.cell(row=6, column=5).value == "日①"


# --- 計算式の答えも一緒に書き込む -------------------------------------------------


def _three_day_schedule():
    """3日ぶん・2名の小さな勤務表。集計の答えが手で数えられる大きさにする。"""
    from datetime import date

    from shift_scheduler.calendar_utils import Day

    days = [Day(date(2026, 10, number)) for number in (1, 2, 3)]
    profiles = [
        StaffProfile(staff_id="a", name="あ", role="介護士", unit="ばら", sheet_label="1"),
        StaffProfile(staff_id="b", name="い", role="看護師", unit="ばら", sheet_label="2"),
    ]
    result = ScheduleResult(status="最適", year=2026, month=10, days=days)
    result.assignments = {
        "a": {1: "日①", 2: "○", 3: "△"},
        "b": {1: "公", 2: "日①", 3: "◉"},
    }
    return result, profiles


def test_集計欄は計算し直さなくても数字が読める(tmp_path):
    """Macの「Numbers」など、開いても計算し直さないアプリのための備え。

    施設の担当者から「日別チェックと集計欄が空欄になる」と報告があった。
    openpyxl は式しか書かないので、式を計算するアプリでしか数字が見えない。
    data_only=True は「保存されている答え」だけを読むので、
    計算し直さないアプリと同じ見え方になる。
    """
    import openpyxl

    from shift_scheduler.excel_export import export_schedule

    result, profiles = _three_day_schedule()
    path = tmp_path / "勤務表.xlsx"
    export_schedule(result, profiles, path)

    sheet = openpyxl.load_workbook(path, data_only=True).active
    summary_column = {label: 8 + offset for offset, label in enumerate(SUMMARY)}

    # 「あ」は 日① / ○ / △。日勤1回・夜勤1回。△は集計に数えない。
    assert sheet.cell(row=6, column=summary_column["日勤"]).value == 1
    assert sheet.cell(row=6, column=summary_column["夜勤"]).value == 1
    assert sheet.cell(row=6, column=summary_column["公休"]).value == 0
    # 「い」は 公 / 日① / ◉。
    assert sheet.cell(row=8, column=summary_column["公休"]).value == 1
    assert sheet.cell(row=8, column=summary_column["深夜"]).value == 1


def test_日別チェックも計算し直さなくても数字が読める(tmp_path):
    import openpyxl

    from shift_scheduler.excel_export import export_schedule

    result, profiles = _three_day_schedule()
    path = tmp_path / "勤務表.xlsx"
    export_schedule(result, profiles, path)

    sheet = openpyxl.load_workbook(path, data_only=True).active
    rows = {
        sheet.cell(row=row, column=4).value: row
        for row in range(1, sheet.max_row + 1)
    }
    # 日① は1日と2日に1人ずつ、3日は0人。
    assert [sheet.cell(row=rows["日①"], column=5 + i).value for i in range(3)] == [1, 1, 0]
    assert [sheet.cell(row=rows["◉"], column=5 + i).value for i in range(3)] == [0, 0, 1]

    judge = next(
        row
        for row in range(1, sheet.max_row + 1)
        if sheet.cell(row=row, column=2).value == "6種そろっているか"
    )
    # 6種のうち日①しか埋まっていないので、どの日も×。
    assert [sheet.cell(row=judge, column=5 + i).value for i in range(3)] == ["×", "×", "×"]


def test_答えを書き込んでも式は残る(tmp_path):
    """Excelで手直ししたときに、その場で計算し直される必要がある。"""
    import openpyxl

    from shift_scheduler.excel_export import export_schedule

    result, profiles = _three_day_schedule()
    path = tmp_path / "勤務表.xlsx"
    export_schedule(result, profiles, path)

    sheet = openpyxl.load_workbook(path).active  # data_only=False で式を読む
    assert sheet.cell(row=6, column=9).value == '=COUNTIF(E6:G6,"○")'
