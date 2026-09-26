"""勤務表をExcelに書き出す。施設で実際に使われている勤務計画表と同じ体裁にする。

  A列  ユニット名(区画の先頭行にだけ入れる)
  B列  職種・肩書き
  C列  勤務表の番号
  D列  日付/曜日/行事 の見出し
  E〜  1日〜末日
  右端 日勤 / 夜勤 / 深夜 / 公休 / 残業 / 有休 / 夏正 / 出研 の集計

1人につき2行。1行目が勤務欄、2行目は中抜けの2コマ目と「せ」(責任者)を書く欄。
最下部に日別人数と、番号付き6種がそろっているかの判定を置く。
"""

from __future__ import annotations

import os
import re
import zipfile
from pathlib import Path
from typing import Dict, List, Sequence
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .calendar_utils import month_label
from .facility import facility
from .models import StaffProfile
from .scheduler import ScheduleResult
from .shifts import (
    ABSENCE_MARKS,
    DAY_SHIFTS,
    HOUR_CHOICES,
    HEALTH_CHECK,
    LATE_NIGHT_IN,
    NIGHT_AFTER,
    NIGHT_IN,
    OFF,
    PAID_LEAVE,
    REQUIRED_DAY_SHIFTS,
    SUMMER_LEAVE,
    TRAINING,
)

SHEET_NAME = "勤務計画表"
RESPONSIBLE = "せ"

UNIT_COLUMN, ROLE_COLUMN, NUMBER_COLUMN, LABEL_COLUMN = 1, 2, 3, 4
FIRST_DAY_COLUMN = 5
HEADER_ROW, WEEKDAY_ROW, EVENT_ROW = 3, 4, 5
FIRST_STAFF_ROW = 6

SUMMARY = ("日勤", "夜勤", "深夜", "公休", "残業", "有休", "夏正", "出研")
# 集計の中身は数字ではなく Excel の計算式で入れる。
# この勤務表はたたき台で、受け取った側が手直しする前提のため、
# セルを書き換えたら右端と下の集計もその場で変わる必要がある。
#
# 「残業」は実物の勤務表にある手書き欄。記号からは決められないので空欄のまま。
# 「日勤」は番号付きの勤務だけを数える。時間を直接書く方のセル(「9-16時」など)は
# 数えない(施設の担当者に確認済み)。
SUMMARY_MARKS: Dict[str, tuple] = {
    "日勤": tuple(DAY_SHIFTS),
    "夜勤": (NIGHT_IN,),
    "深夜": (LATE_NIGHT_IN,),
    "公休": (OFF,),
    "残業": (),
    "有休": (PAID_LEAVE,),
    "夏正": (SUMMER_LEAVE,),
    "出研": (TRAINING, HEALTH_CHECK),
}
UNIT_ORDER = facility().units

# 手直しするときにセルのプルダウンに出す記号。
# ◉ や 丸数字 は手で打ちにくく、○ は似た字(◯ 〇)と取り違えやすい。
# 見た目が同じでも別の文字だと集計の式が拾わないので、選べるようにする。
# ただし入力を禁止はしない(_add_mark_dropdown を参照)。
EDIT_CHOICES = (
    (OFF,)
    + tuple(DAY_SHIFTS)
    + HOUR_CHOICES
    + (NIGHT_IN, NIGHT_AFTER, LATE_NIGHT_IN)
    + ABSENCE_MARKS
)

_HEAD_FILL = PatternFill("solid", fgColor="2F5D8C")
_WEEKEND_FILL = PatternFill("solid", fgColor="FDECEC")
_UNIT_FILL = PatternFill("solid", fgColor="EDF0F4")
_NIGHT_FILL = PatternFill("solid", fgColor="E4D7F5")
_OFF_FILL = PatternFill("solid", fgColor="F2F2F2")
_NG_FILL = PatternFill("solid", fgColor="FBE3E3")
_THIN = Side(style="thin", color="BFC7D1")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

# 記号ごとの色。目で追いやすくするためだけのもの。
_MARK_FILL = {
    NIGHT_IN: _NIGHT_FILL,
    NIGHT_AFTER: PatternFill("solid", fgColor="EFE8FA"),
    LATE_NIGHT_IN: PatternFill("solid", fgColor="D9D2E9"),
    OFF: _OFF_FILL,
    PAID_LEAVE: PatternFill("solid", fgColor="FFF2CC"),
    SUMMER_LEAVE: PatternFill("solid", fgColor="FFF2CC"),
    TRAINING: PatternFill("solid", fgColor="E2EFDA"),
    HEALTH_CHECK: PatternFill("solid", fgColor="E2EFDA"),
}


def export_schedule(
    result: ScheduleResult,
    profiles: Sequence[StaffProfile],
    path: Path,
    events: Dict[int, str] = None,
) -> None:
    """組み上がった勤務表をExcelにする。events は {日にち: 行事名}。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    days = result.days
    events = events or {}

    _write_legend_row(sheet, len(days))
    _write_title(sheet, result)
    _write_header(sheet, days, events)

    # 計算式のセルについて、その答えをここに控えておく。
    # {"AJ6": 20} のように、セルの位置と答えを入れる。保存の後で書き込む。
    cached: Dict[str, object] = {}

    ordered = _order_staff(profiles, result)
    row = FIRST_STAFF_ROW
    previous_unit = None
    for profile in ordered:
        unit = _unit_label(profile)
        _write_staff(
            sheet, row, profile, result, unit if unit != previous_unit else "", cached
        )
        previous_unit = unit
        row += 2

    staff_rows = (FIRST_STAFF_ROW, max(FIRST_STAFF_ROW, row - 1))
    _write_daily_check(sheet, row + 1, days, staff_rows, ordered, result, cached)
    _add_mark_dropdown(
        sheet,
        [FIRST_STAFF_ROW + index * 2 for index in range(len(ordered))],
        len(days),
    )
    _finish_layout(sheet, days, len(ordered))

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    _write_cached_values(path, cached)


# --- 見出し -------------------------------------------------------------------


def _write_legend_row(sheet, day_count: int) -> None:
    """1行目にシフトの時間を並べる(実物と同じ)。"""
    column = FIRST_DAY_COLUMN
    for code, (start, end) in DAY_SHIFTS.items():
        cell = sheet.cell(row=1, column=column, value=f"{code} {start}-{end}")
        cell.font = Font(size=9)
        column += 4


def _write_title(sheet, result: ScheduleResult) -> None:
    sheet.cell(row=2, column=1, value=facility().ward_label).font = Font(bold=True, size=12)
    sheet.cell(row=2, column=FIRST_DAY_COLUMN, value=month_label(result.year, result.month))


def _write_header(sheet, days, events: Dict[int, str]) -> None:
    for column, label in ((ROLE_COLUMN, "職種"), (LABEL_COLUMN, "日付")):
        _head(sheet, HEADER_ROW, column, label)
    _head(sheet, WEEKDAY_ROW, LABEL_COLUMN, "曜日")
    _head(sheet, EVENT_ROW, LABEL_COLUMN, "行事")

    for index, day in enumerate(days):
        column = FIRST_DAY_COLUMN + index
        _head(sheet, HEADER_ROW, column, day.day)
        _head(sheet, WEEKDAY_ROW, column, day.weekday_name)
        event = sheet.cell(row=EVENT_ROW, column=column, value=events.get(day.day))
        event.alignment = Alignment(horizontal="center", wrap_text=True)
        event.border = _BORDER
        event.font = Font(size=8)
        if day.is_weekend:
            event.fill = _WEEKEND_FILL

    for offset, label in enumerate(SUMMARY):
        _head(sheet, HEADER_ROW, FIRST_DAY_COLUMN + len(days) + offset, label)


def _head(sheet, row: int, column: int, value) -> None:
    cell = sheet.cell(row=row, column=column, value=value)
    cell.font = Font(bold=True, color="FFFFFF", size=10)
    cell.fill = _HEAD_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = _BORDER


# --- スタッフ行 ----------------------------------------------------------------


def _order_staff(profiles, result) -> List[StaffProfile]:
    """師長 → ユニット順 → 介護補助。勤務表に載らない人は除く。"""
    listed = [p for p in profiles if p.staff_id in result.assignments]

    def key(profile):
        if profile.is_head_nurse:
            return (0, "")
        if profile.is_support_staff:
            return (5, profile.sheet_label)
        rank = UNIT_ORDER.index(profile.unit) + 1 if profile.unit in UNIT_ORDER else 6
        return (rank, _number_key(profile.sheet_label))

    return sorted(listed, key=key)


def _number_key(label: str):
    """看護(1,2,3…)を先に、介護(①②③…)を後に並べる。"""
    circles = "①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳㉑㉒㉓㉔㉕㉖"
    if label and label[0] in circles:
        # 丸数字は isdigit() でも真になるので、先に判定する
        return (1, circles.index(label[0]))
    if label.isdigit():
        return (0, int(label))
    return (2, label)


def _unit_label(profile: StaffProfile) -> str:
    if profile.is_head_nurse:
        return ""
    if profile.is_support_staff:
        return "介護補助"
    return profile.unit


def _write_staff(sheet, row: int, profile, result, unit_label: str, cached) -> None:
    if unit_label:
        cell = sheet.cell(row=row, column=UNIT_COLUMN, value=unit_label)
        cell.font = Font(bold=True, size=10)
        cell.fill = _UNIT_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")

    role = sheet.cell(row=row, column=ROLE_COLUMN, value=_role_label(profile))
    role.font = Font(size=9)
    role.alignment = Alignment(vertical="center", wrap_text=True)
    role.border = _BORDER

    number = sheet.cell(row=row, column=NUMBER_COLUMN, value=profile.sheet_label)
    number.alignment = Alignment(horizontal="center", vertical="center")
    number.border = _BORDER

    sheet.cell(row=row, column=LABEL_COLUMN, value=profile.name).border = _BORDER
    sheet.cell(row=row + 1, column=LABEL_COLUMN).border = _BORDER

    assignment = result.assignments.get(profile.staff_id, {})

    for index, day in enumerate(result.days):
        column = FIRST_DAY_COLUMN + index
        mark = assignment.get(day.day, "")
        main, second = _split_mark(mark)

        _write_mark(sheet, row, column, main, day.is_weekend)
        extra = second
        if result.responsible.get(day.day) == profile.staff_id:
            extra = f"{extra} {RESPONSIBLE}".strip() if extra else RESPONSIBLE
        _write_mark(sheet, row + 1, column, extra, day.is_weekend, small=True)

    days_range = _row_range(row, len(result.days))
    # 集計の対象は1行目の勤務欄だけ。2行目(中抜けの2コマ目と「せ」)は数えない。
    marks_in_row = [_split_mark(assignment.get(day.day, ""))[0] for day in result.days]
    for offset, label in enumerate(SUMMARY):
        column = FIRST_DAY_COLUMN + len(result.days) + offset
        marks = SUMMARY_MARKS[label]
        cell = sheet.cell(
            row=row, column=column, value=_countif_formula(days_range, marks)
        )
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _BORDER
        if marks:
            cached[f"{get_column_letter(column)}{row}"] = sum(
                1 for mark in marks_in_row if mark in marks
            )


def _role_label(profile: StaffProfile) -> str:
    if profile.is_head_nurse:
        return "師長"
    return profile.role or ""


def _split_mark(mark: str):
    """中抜け(「09:00-12:00・19:00-20:00」)は2行に分ける。"""
    if mark and "・" in mark:
        first, _, rest = mark.partition("・")
        return first, rest
    return mark, ""


def _write_mark(sheet, row: int, column: int, value: str, weekend: bool, small=False) -> None:
    cell = sheet.cell(row=row, column=column, value=value or None)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = _BORDER
    cell.font = Font(size=8 if small else 10)
    if value in _MARK_FILL:
        cell.fill = _MARK_FILL[value]
    elif weekend:
        cell.fill = _WEEKEND_FILL


def _row_range(row: int, day_count: int) -> str:
    """その人の勤務欄(1行分)の範囲。例: E6:AI6"""
    first = get_column_letter(FIRST_DAY_COLUMN)
    last = get_column_letter(FIRST_DAY_COLUMN + day_count - 1)
    return f"{first}{row}:{last}{row}"


def _countif_formula(cell_range: str, marks) -> str:
    """その範囲に marks がいくつあるかを数える式。marks が空なら空欄。

    数字ではなく式で入れるのは、受け取った側が勤務表を手直ししたときに
    集計がその場で変わるようにするため。
    """
    if not marks:
        return None
    return "=" + "+".join(f'COUNTIF({cell_range},"{mark}")' for mark in marks)


# --- 計算式の答えを書き込む ------------------------------------------------------


def _write_cached_values(path: Path, cached: Dict[str, object]) -> None:
    """計算式のセルに、その答えも一緒に書き込む。

    Excelのファイルは、計算式のセルに「式」と「前回の答え」の両方を持てる。
    Excelは開いたときに式を計算し直すが、Macの「Numbers」など他のアプリは
    計算し直さず、保存されている答えをそのまま見せる。
    openpyxl は式しか書かないので、そういうアプリでは集計欄が
    まるごと空欄に見えてしまう(施設の担当者から報告あり)。

    そこで、保存した後のファイルを開いて、式のとなりに答えを差し込む。
    式は残したままなので、Excelで手直しすればその場で計算し直される。
    """
    if not cached:
        return

    # この勤務表はシート1枚だけ(export_schedule が1枚しか作らない)。
    # シートを増やすときは、ここも直すこと。
    sheet_name = "xl/worksheets/sheet1.xml"
    with zipfile.ZipFile(path) as source:
        names = source.namelist()
        if sheet_name not in names:
            return  # 想定しない構成。何もしないでおく(壊すよりは空欄のまま)
        contents = {name: source.read(name) for name in names}

    contents[sheet_name] = _insert_values(
        contents[sheet_name].decode("utf-8"), cached
    ).encode("utf-8")

    # いったん別の名前で書いてから置き換える。
    # 同じファイルに直接上書きすると、書いている途中で止まったときに
    # お客様にお渡しするファイルが壊れたまま残ってしまう。
    temporary = path.with_name(path.name + ".tmp")
    try:
        with zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as destination:
            for name in names:
                destination.writestr(name, contents[name])
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()


# 計算式が入っているセル1つ分。
# openpyxl は答えの場所を <v /> と空で書くので、そこも一緒に拾って置き換える。
#   <c r="AJ6" s="5"><f>COUNTIF(...)</f><v /></c>
_FORMULA_CELL = re.compile(
    r'<c r="([A-Z]+\d+)"([^>]*)>(<f[^>]*>.*?</f>)(?:<v\s*/>|<v>.*?</v>)?</c>',
    re.DOTALL,
)


def _insert_values(xml: str, cached: Dict[str, object]) -> str:
    """セルの式のうしろに <v>答え</v> を差し込む。"""

    def replace(match):
        reference, attributes, formula = match.groups()
        if reference not in cached:
            return match.group(0)
        value = cached[reference]
        # 型の指定は書き直す。文字が答えのときは t="str" が要る。
        attributes = re.sub(r'\s*t="[^"]*"', "", attributes)
        if isinstance(value, str):
            attributes += ' t="str"'
        return f'<c r="{reference}"{attributes}>{formula}<v>{escape(str(value))}</v></c>'

    return _FORMULA_CELL.sub(replace, xml)


# --- 最下部の集計 --------------------------------------------------------------


def _write_daily_check(sheet, row: int, days, staff_rows, ordered, result, cached) -> None:
    """番号付き6種が毎日そろっているかを確かめる行。実物にも同じものがある。

    こちらも数字ではなく式で入れる。勤務表を手直ししたときに、
    その日がそろっているかどうかがその場で分かるようにするため。

    数える範囲はスタッフ欄ぜんぶ(1人2行の2行目も含む)。2行目には中抜けの
    時間と「せ」しか入らず、番号付きの記号とは一致しないので混ざらない。
    """
    label = sheet.cell(row=row, column=ROLE_COLUMN, value="日別チェック")
    label.font = Font(bold=True, size=9)

    first_staff_row, last_staff_row = staff_rows
    codes = REQUIRED_DAY_SHIFTS + (NIGHT_IN, LATE_NIGHT_IN)
    count_rows = {}
    # その日に誰がどの記号だったか。式の答えを控えるために数える。
    # 集計と同じく、数えるのは1行目の勤務欄だけ。
    marks_on_day = {
        day.day: [
            _split_mark(result.assignments.get(profile.staff_id, {}).get(day.day, ""))[0]
            for profile in ordered
        ]
        for day in days
    }

    for offset, code in enumerate(codes):
        line = row + 1 + offset
        count_rows[code] = line
        sheet.cell(row=line, column=LABEL_COLUMN, value=code).alignment = Alignment(
            horizontal="center"
        )
        for index, day in enumerate(days):
            letter = get_column_letter(FIRST_DAY_COLUMN + index)
            cell = sheet.cell(
                row=line,
                column=FIRST_DAY_COLUMN + index,
                value=f'=COUNTIF({letter}{first_staff_row}:{letter}{last_staff_row},"{code}")',
            )
            cell.alignment = Alignment(horizontal="center")
            cell.font = Font(size=8)
            cell.border = _BORDER
            cached[f"{letter}{line}"] = marks_on_day[day.day].count(code)

    judge = row + 1 + len(codes) + 1
    sheet.cell(row=judge, column=ROLE_COLUMN, value="6種そろっているか").font = Font(
        bold=True, size=9
    )
    for index, day in enumerate(days):
        letter = get_column_letter(FIRST_DAY_COLUMN + index)
        conditions = ",".join(
            f"{letter}{count_rows[code]}=1" for code in REQUIRED_DAY_SHIFTS
        )
        cell = sheet.cell(
            row=judge,
            column=FIRST_DAY_COLUMN + index,
            value=f'=IF(AND({conditions}),"○","×")',
        )
        cell.alignment = Alignment(horizontal="center")
        cell.border = _BORDER
        complete = all(
            marks_on_day[day.day].count(code) == 1 for code in REQUIRED_DAY_SHIFTS
        )
        cached[f"{letter}{judge}"] = "○" if complete else "×"

    _mark_shortage(sheet, judge, len(days))


def _mark_shortage(sheet, judge_row: int, day_count: int) -> None:
    """そろっていない日に色を付ける。

    色も式で決める(条件付き書式)。固定の色だと、手直ししてそろった後も
    赤いままになってしまう。
    """
    if not day_count:
        return
    first = get_column_letter(FIRST_DAY_COLUMN)
    last = get_column_letter(FIRST_DAY_COLUMN + day_count - 1)
    sheet.conditional_formatting.add(
        f"{first}{judge_row}:{last}{judge_row}",
        CellIsRule(operator="equal", formula=['"×"'], fill=_NG_FILL),
    )


def _add_mark_dropdown(sheet, staff_rows, day_count: int) -> None:
    """勤務欄のセルに、記号を選ぶプルダウンを付ける。

    手直しするときに ◉ や 丸数字 を打つ手間をなくし、
    ○ と ◯ の取り違えで集計が合わなくなるのを防ぐ。

    一覧に無いものを打ってもエラーにしない(showErrorMessage=False)。
    勤務時間を直接書く方のセルには「15」「19-20」などが入るうえ、
    こちらが想定しない直し方をされることもあるため、入力は禁止しない。

    付けるのは1人につき1行目の勤務欄だけ。2行目は中抜けと「せ」の欄で
    用途が違い、右端と最下部は計算式なので付けない。
    """
    if not staff_rows or not day_count:
        return

    # 記号はカンマで区切って並べる。カンマを含む記号は足さないこと。
    options = f'"{",".join(EDIT_CHOICES)}"'
    if len(options) > 255:  # Excelの上限。前後の " も数に入る。超えると壊れる
        return

    validation = DataValidation(
        type="list",
        formula1=options,
        allow_blank=True,
        showErrorMessage=False,
    )
    sheet.add_data_validation(validation)

    first = get_column_letter(FIRST_DAY_COLUMN)
    last = get_column_letter(FIRST_DAY_COLUMN + day_count - 1)
    for row in staff_rows:
        validation.add(f"{first}{row}:{last}{row}")


def _finish_layout(sheet, days, staff_count: int) -> None:
    sheet.column_dimensions["A"].width = 9
    sheet.column_dimensions["B"].width = 16
    sheet.column_dimensions["C"].width = 5
    sheet.column_dimensions["D"].width = 9
    for index in range(len(days)):
        sheet.column_dimensions[get_column_letter(FIRST_DAY_COLUMN + index)].width = 4.5
    for offset in range(len(SUMMARY)):
        sheet.column_dimensions[
            get_column_letter(FIRST_DAY_COLUMN + len(days) + offset)
        ].width = 5.5
    sheet.freeze_panes = sheet.cell(row=FIRST_STAFF_ROW, column=FIRST_DAY_COLUMN)
